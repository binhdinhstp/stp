import os,re,sqlite3,secrets,hashlib,hmac,io,calendar
from functools import wraps
from datetime import datetime
from pathlib import Path
from flask import Flask,request,session,redirect,url_for,render_template,flash,send_file,abort
from werkzeug.security import generate_password_hash,check_password_hash
from docx import Document
from docx.shared import Pt

BASE=Path(__file__).resolve().parent
app=Flask(__name__)
app.config['SECRET_KEY']=os.environ.get('SECRET_KEY',secrets.token_hex(32))
app.config.update(SESSION_COOKIE_HTTPONLY=True,SESSION_COOKIE_SAMESITE='Lax',MAX_CONTENT_LENGTH=2*1024*1024)
if os.environ.get('HTTPS_ONLY')=='1':app.config['SESSION_COOKIE_SECURE']=True
DB=Path(os.environ.get('DATABASE_PATH',str(BASE/'reports.sqlite3')))
TEMPLATE=BASE/'mau_bao_cao_thang.docx'

def db():
 c=sqlite3.connect(DB,timeout=15);c.row_factory=sqlite3.Row;c.execute('PRAGMA foreign_keys=ON');return c

def init():
 DB.parent.mkdir(parents=True,exist_ok=True)
 with db() as c:
  c.execute('PRAGMA journal_mode=WAL')
  c.executescript('''CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY,username TEXT UNIQUE NOT NULL,password TEXT NOT NULL,department TEXT NOT NULL,role TEXT NOT NULL,active INTEGER DEFAULT 1);
CREATE TABLE IF NOT EXISTS entries(id INTEGER PRIMARY KEY,period TEXT NOT NULL,kind TEXT NOT NULL,section TEXT NOT NULL,department TEXT NOT NULL,content TEXT NOT NULL,spell_hash TEXT NOT NULL,reviewed_by INTEGER NOT NULL,reviewed_at TEXT NOT NULL,updated_by INTEGER NOT NULL,updated_at TEXT NOT NULL,UNIQUE(period,kind,section,department));
CREATE TABLE IF NOT EXISTS audits(id INTEGER PRIMARY KEY,user_id INTEGER,action TEXT,detail TEXT,created_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS locks(period TEXT PRIMARY KEY,locked_by INTEGER,locked_at TEXT);''')
  if not c.execute('SELECT 1 FROM users LIMIT 1').fetchone():
   pwd=os.environ.get('ADMIN_PASSWORD','')
   if not pwd:raise RuntimeError('Cần thiết lập biến môi trường ADMIN_PASSWORD trước lần chạy đầu tiên.')
   c.execute('INSERT INTO users(username,password,department,role) VALUES(?,?,?,?)',('admin',generate_password_hash(pwd),'Văn phòng Sở','admin'))
init()

def paragraphs():
 d=Document(TEMPLATE)
 return [(i,p.text.strip()) for i,p in enumerate(d.paragraphs)]
PARAS=paragraphs()
# Extract headings directly from uploaded model, retaining spelling and order.
GROUP_INDICES=[9,20,30,37,52,54,56]
CHILD_INDICES=[[10,12,14,16,18],[21,23,26,28],[31,33,35],[38,40,42,44,46,48,50],[52],[54],[56]]
PLAN_INDICES=[59,61,63,65,67,69]
HEAD={i:dict(PARAS).get(i,'') for i in set(sum(CHILD_INDICES,[])+GROUP_INDICES+PLAN_INDICES)}
FIELDS=[{'id':f'r{i}','heading':HEAD[i],'group':gi+1,'part':'result','paragraph':i} for gi,ls in enumerate(CHILD_INDICES) for i in ls]+[{'id':f'p{i}','heading':HEAD[i],'group':0,'part':'plan','paragraph':i} for i in PLAN_INDICES]
LOOKUP={x['id']:x for x in FIELDS}
# Small, explicitly bounded rule-based pre-save review; not an AI proofreading claim.
RULES=[
 ('triễn khai','triển khai'),('kiễm tra','kiểm tra'),('thẫm định','thẩm định'),
 ('phỗ biến','phổ biến'),('thưc hiện','thực hiện'),('quy đinh','quy định'),
 ('giãi quyết','giải quyết'),('xữ lý','xử lý'),('hổ trợ','hỗ trợ'),
 ('pháp luât','pháp luật'),('báo cao','báo cáo'),('công tac','công tác'),
 ('tư phap','tư pháp'),('thực hiên','thực hiện'),('đươc','được'),
 ('đia phương','địa phương'),('địa phuơng','địa phương'),
 ('cơ quan nhà nuớc','cơ quan nhà nước'),('quãn lý','quản lý'),
 ('kiên nghị','kiến nghị'),('phản anh','phản ánh'),('văn bãn','văn bản'),
 ('thủ tuc hành chính','thủ tục hành chính'),('hộ tich','hộ tịch'),
 ('công chưng','công chứng'),('đấu gia tài sản','đấu giá tài sản'),
 ('phòng chông tham nhũng','phòng, chống tham nhũng'),
 ('giám đinh tư pháp','giám định tư pháp'),('hòa giãi','hòa giải'),
 ('nghiệp vu','nghiệp vụ'),('kết quã','kết quả'),('cấp xa','cấp xã'),
 ('đánh gia','đánh giá'),('quốc tich','quốc tịch'),
]
def issues(s):
 r=[]
 for old,new in RULES:
  for m in re.finditer(re.escape(old),s,re.I):
   r.append({'at':m.start(),'end':m.end(),'before':m.group(),'after':new,'label':'Từ/cụm từ có khả năng sai chính tả'})
 for m in re.finditer(r' {2,}',s):
  r.append({'at':m.start(),'end':m.end(),'before':m.group(),'after':' ','label':'Khoảng trắng liên tiếp'})
 for m in re.finditer(r' +[,.;:!?]',s):
  r.append({'at':m.start(),'end':m.end(),'before':m.group(),'after':m.group().lstrip(),'label':'Khoảng trắng trước dấu câu'})
 r.sort(key=lambda x:(x['at'],-x['end']))
 # Filter overlapping suggestions so applying one will not corrupt another.
 out=[];last=-1
 for item in r:
  if item['at']>=last:out.append(item);last=item['end']
 return out
def digest(s):return hashlib.sha256(s.encode('utf-8')).hexdigest()
def csrf():return session.setdefault('csrf',secrets.token_urlsafe(24))
app.jinja_env.globals['csrf']=csrf
@app.before_request
def protect():
 if request.method=='POST' and not hmac.compare_digest(str(request.form.get('csrf','')),str(csrf())):abort(400,'Token phiên làm việc không hợp lệ')
@app.after_request
def headers(r):
 r.headers['X-Content-Type-Options']='nosniff';r.headers['X-Frame-Options']='DENY';r.headers['Referrer-Policy']='same-origin';r.headers['Cache-Control']='no-store';return r

def login_required(f):
 @wraps(f)
 def wrapped(*a,**kw):
  if 'uid' not in session:return redirect(url_for('login'))
  return f(*a,**kw)
 return wrapped
def current():
 with db() as c:return c.execute('SELECT * FROM users WHERE id=? AND active=1',(session['uid'],)).fetchone()
def context():
 u=current();return u or abort(403)
def period():
 k=request.values.get('kind','month');yr=request.values.get('year','2026');mo=request.values.get('month','10')
 if k not in ['month','year'] or not re.fullmatch(r'20\d{2}',yr) or not mo.isdigit() or not 1<=int(mo)<=12:abort(400)
 return k,yr,mo,yr+'-'+(mo.zfill(2) if k=='month' else 'year')
def is_locked(p):
 with db() as c:return bool(c.execute('SELECT 1 FROM locks WHERE period=?',(p,)).fetchone())
@app.route('/login',methods=['GET','POST'])
def login():
 if request.method=='POST':
  with db() as c:u=c.execute('SELECT * FROM users WHERE username=? AND active=1',(request.form.get('username',''),)).fetchone()
  if u and check_password_hash(u['password'],request.form.get('password','')):
   session.clear();session['uid']=u['id'];session['csrf']=secrets.token_urlsafe(24);return redirect(url_for('index'))
  flash('Sai thông tin đăng nhập.','error')
 return render_template('login.html')
@app.post('/logout')
def logout():session.clear();return redirect(url_for('login'))
@app.route('/')
@login_required
def index():
 u=context();k,y,m,p=period()
 with db() as c:rows=c.execute('SELECT * FROM entries WHERE period=?',(p,)).fetchall()
 mine={r['section']:r for r in rows if r['department']==u['department']}
 counts={f['id']:len(set(r['department'] for r in rows if r['section']==f['id'])) for f in FIELDS}
 return render_template('index.html',user=u,kind=k,year=y,month=m,period=p,groups=[(HEAD[i],[x for x in FIELDS if x['group']==gi+1 and x['part']=='result']) for gi,i in enumerate(GROUP_INDICES)],plans=[x for x in FIELDS if x['part']=='plan'],mine=mine,counts=counts,locked=is_locked(p))
@app.route('/entry/<sid>',methods=['GET','POST'])
@login_required
def entry(sid):
 u=context();f=LOOKUP.get(sid)
 if not f:abort(404)
 k,y,m,p=period()
 with db() as c:r=c.execute('SELECT * FROM entries WHERE period=? AND section=? AND department=?',(p,sid,u['department'])).fetchone()
 if request.method=='POST':
  if is_locked(p):abort(403,'Kỳ báo cáo đã khóa')
  content=request.form.get('content','').strip()
  if len(content)>50000:abort(400,'Nội dung vượt 50.000 ký tự')
  action=request.form.get('action','check');findings=issues(content)
  if action=='save':
   if not content:flash('Không thể lưu nội dung rỗng.','error')
   elif findings:flash('Còn lỗi được bộ quy tắc phát hiện; hãy sửa và kiểm tra lại.','error')
   elif not request.form.get('ack'):flash('Bạn phải xác nhận đã kiểm tra nội dung.','error')
   elif not hmac.compare_digest(request.form.get('proof',''),hmac.new(app.secret_key.encode(),(str(u['id'])+'|'+p+'|'+sid+'|'+digest(content)).encode(),hashlib.sha256).hexdigest()):flash('Nội dung đã thay đổi; phải kiểm tra lại trước khi lưu.','error')
   else:
    now=datetime.now().isoformat(timespec='seconds')
    with db() as c:
     c.execute('INSERT INTO entries(period,kind,section,department,content,spell_hash,reviewed_by,reviewed_at,updated_by,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?) ON CONFLICT(period,kind,section,department) DO UPDATE SET content=excluded.content,spell_hash=excluded.spell_hash,reviewed_by=excluded.reviewed_by,reviewed_at=excluded.reviewed_at,updated_by=excluded.updated_by,updated_at=excluded.updated_at',(p,k,sid,u['department'],content,digest(content),u['id'],now,u['id'],now))
     c.execute('INSERT INTO audits(user_id,action,detail,created_at) VALUES(?,?,?,?)',(u['id'],'save',p+' / '+sid+' / '+u['department'],now))
    flash('Đã lưu nội dung sau khi xác nhận kiểm tra.','success');return redirect(url_for('index',kind=k,year=y,month=m))
  if not content:findings=[]
  proof=hmac.new(app.secret_key.encode(),(str(u['id'])+'|'+p+'|'+sid+'|'+digest(content)).encode(),hashlib.sha256).hexdigest() if content else ''
  return render_template('entry.html',user=u,field=f,kind=k,year=y,month=m,period=p,content=content,findings=findings,checked=bool(content),proof=proof,locked=is_locked(p))
 return render_template('entry.html',user=u,field=f,kind=k,year=y,month=m,period=p,content=r['content'] if r else '',findings=[],checked=False,proof='',locked=is_locked(p))
@app.route('/report')
@login_required
def report():
 u=context();k,y,m,p=period()
 with db() as c:rs=c.execute('SELECT * FROM entries WHERE period=? ORDER BY department',(p,)).fetchall()
 rows={f['id']:[r for r in rs if r['section']==f['id'] and (u['role']=='admin' or r['department']==u['department'])] for f in FIELDS}
 return render_template('report.html',user=u,kind=k,year=y,month=m,period=p,groups=[(HEAD[i],[x for x in FIELDS if x['group']==gi+1 and x['part']=='result']) for gi,i in enumerate(GROUP_INDICES)],plans=[x for x in FIELDS if x['part']=='plan'],rows=rows)
@app.get('/export')
@login_required
def export():
 u=context();k,y,m,p=period()
 if k!='month':abort(400,'Chưa có mẫu Word báo cáo năm chính thức. Vui lòng cung cấp mẫu năm.')
 if u['role']!='admin':abort(403,'Chỉ Văn phòng Sở được xuất báo cáo tổng hợp')
 with db() as c:rs=c.execute('SELECT * FROM entries WHERE period=? ORDER BY department',(p,)).fetchall()
 d=Document(TEMPLATE)
 nextmonth=int(m)%12+1;nextyear=int(y)+(int(m)==12)
 for para in d.paragraphs:
  t=para.text
  if 'Kết quả thực hiện công tác tư pháp tháng' in t and 'và chương trình' not in t:para.text=f'Kết quả thực hiện công tác tư pháp tháng {int(m)}/{y}'
  elif 'và chương trình công tác tư pháp tháng' in t:para.text=f'và chương trình công tác tư pháp tháng {nextmonth}/{nextyear}'
  elif t.startswith('II. CHƯƠNG TRÌNH CÔNG TÁC THÁNG'):para.text=f'II. CHƯƠNG TRÌNH CÔNG TÁC THÁNG {nextmonth} NĂM {nextyear}'
  elif t.startswith('Trên đây là báo cáo kết quả'):para.text=f'Trên đây là báo cáo kết quả thực hiện công tác tư pháp tháng {int(m)}/{y} và chương trình công tác tháng {nextmonth}/{nextyear} của Sở Tư pháp./.'
 # Use the source document's original paragraph positions, avoiding ambiguity of duplicate headings.
 for f in FIELDS:
  idx=f['paragraph']
  # If parent heading has a content line, or heading has two nested TGPL sections, next paragraph is correct.
  target=idx+1
  matches=[r for r in rs if r['section']==f['id']]
  if target>=len(d.paragraphs):continue
  value='\n'.join('['+r['department']+'] '+r['content'] for r in matches) if matches else 'Nội dung các phòng, đơn vị cập nhật'
  para=d.paragraphs[target];para.text=value
  for run in para.runs:run.font.name='Times New Roman';run.font.size=Pt(13)
 out=io.BytesIO();d.save(out);out.seek(0)
 return send_file(out,mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',as_attachment=True,download_name=f'Bao_cao_tu_phap_thang_{int(m)}_{y}.docx')
@app.route('/admin',methods=['GET','POST'])
@login_required
def admin():
 u=context()
 if u['role']!='admin':abort(403)
 if request.method=='POST':
  action=request.form.get('action')
  if action=='create':
   name=request.form.get('username','').strip();pw=request.form.get('password','');dept=request.form.get('department','').strip();role=request.form.get('role','editor')
   if not re.fullmatch(r'[a-zA-Z0-9_.-]{3,40}',name) or len(pw)<12 or not dept or role not in ('editor','admin'):flash('Dữ liệu không hợp lệ. Mật khẩu ít nhất 12 ký tự.','error')
   else:
    try:
     with db() as c:c.execute('INSERT INTO users(username,password,department,role) VALUES(?,?,?,?)',(name,generate_password_hash(pw),dept,role))
     flash('Đã tạo tài khoản.','success')
    except sqlite3.IntegrityError:flash('Tên đăng nhập đã tồn tại.','error')
  elif action=='lock':
   _,_,_,p=period()
   with db() as c:c.execute('INSERT OR REPLACE INTO locks(period,locked_by,locked_at) VALUES(?,?,?)',(p,u['id'],datetime.now().isoformat()))
   flash('Đã khóa kỳ báo cáo.','success')
  elif action=='unlock':
   _,_,_,p=period()
   with db() as c:c.execute('DELETE FROM locks WHERE period=?',(p,))
   flash('Đã mở khóa kỳ báo cáo.','success')
 with db() as c:users=c.execute('SELECT id,username,department,role,active FROM users ORDER BY id').fetchall()
 return render_template('admin.html',user=u,users=users)
@app.route('/change-password',methods=['GET','POST'])
@login_required
def change_password():
 u=context()
 if request.method=='POST':
  if not check_password_hash(u['password'],request.form.get('old','')):flash('Mật khẩu cũ không đúng.','error')
  elif len(request.form.get('new',''))<12:flash('Mật khẩu mới phải từ 12 ký tự.','error')
  else:
   with db() as c:c.execute('UPDATE users SET password=? WHERE id=?',(generate_password_hash(request.form['new']),u['id']))
   flash('Đã đổi mật khẩu.','success');return redirect(url_for('index'))
 return render_template('password.html',user=u)
if __name__=='__main__':app.run(host='127.0.0.1',port=int(os.environ.get('PORT','8080')),debug=False)
